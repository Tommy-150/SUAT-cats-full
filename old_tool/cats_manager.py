#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SUAT-cats 管理器（增强版：标签管理、批量打标签、窗口自适应）
========================================================
行列表视图，支持上移/下移步数调整顺序，内建预览服务器。
集成：基本信息编辑、故事编辑、图片预览、缩略图编辑、新增猫咪、删除猫咪、
     标签管理（增删改、查看关联猫、批量添加标签）、敏感词管理、自动同步前端（Excel + cats.json）。
关闭时若有未保存的顺序变更会提醒。
依赖：Pillow、openpyxl
"""

import os
import re
import sys
import json
import glob
import shutil
import traceback
import subprocess
import threading
import webbrowser
from pathlib import Path
from http.server import HTTPServer, SimpleHTTPRequestHandler

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, colorchooser
import tkinter.font as tkfont

try:
    from PIL import Image, ImageTk
except ImportError:
    print("缺少 Pillow：pip install Pillow")
    sys.exit(1)
try:
    from openpyxl import load_workbook
except ImportError:
    print("缺少 openpyxl：pip install openpyxl")
    sys.exit(1)


# ============================================================
# Windows 高 DPI 适配
# ============================================================
def _enable_dpi_awareness():
    if sys.platform != "win32":
        return
    try:
        import ctypes
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(2)
        except Exception:
            try:
                ctypes.windll.user32.SetProcessDPIAware()
            except Exception:
                pass
    except Exception:
        pass


_enable_dpi_awareness()


# ============================================================
# 视觉常量
# ============================================================
COLOR_BG = "#F7F5F0"
COLOR_CARD = "#FFFFFF"
COLOR_CARD_BORDER = "#ECE7DC"
COLOR_TEXT = "#333333"
COLOR_SUBTEXT = "#888888"
COLOR_ACCENT = "#4A4A4A"
COLOR_ACCENT_LIGHT = "#6c6c6c"
COLOR_MALE = "#5B8DEF"
COLOR_FEMALE = "#FF7E93"
COLOR_UNKNOWN = "#AAAAAA"
COLOR_PAW = "#D9A05B"
COLOR_HOVER = "#FAF7EE"
COLOR_DIVIDER = "#EEE9DD"
COLOR_DANGER = "#E57373"
COLOR_OK = "#7CB87C"
COLOR_WARN = "#E8A53A"
COLOR_INFO = "#5B8DEF"

FONT_CANDIDATES_UI = [
    "Noto Sans CJK SC", "Noto Sans SC", "Source Han Sans SC",
    "Source Han Sans CN", "思源黑体", "思源黑体 CN",
    "Sarasa Gothic SC", "Sarasa UI SC",
    "LXGW WenKai", "LXGW WenKai GB", "霞鹜文楷",
    "WenQuanYi Micro Hei", "文泉驿微米黑",
    "DejaVu Sans", "Liberation Sans",
]
FONT_CANDIDATES_MONO = [
    "JetBrains Mono", "Sarasa Mono SC", "Source Code Pro",
    "DejaVu Sans Mono", "Liberation Mono",
]

FONT_FAMILY_UI = None
FONT_FAMILY_MONO = None

FS_TITLE = 22
FS_SUBTITLE = 11
FS_CARD_NAME = 13
FS_CARD_DESC = 9
FS_BODY = 10
FS_SMALL = 9
FS_TINY = 8


# ============================================================
# 路径与字段
# ============================================================
BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = BASE_DIR / "统计信息.xlsx"
JSON_PATH = BASE_DIR / "cats.json"
TAG_COLOR_PATH = BASE_DIR / "tag_color.json"
CLASSIFIED_DIR = BASE_DIR / "classified"
INDEX_HTML = BASE_DIR / "index.html"
TEMP_THUMB_DIR = BASE_DIR / ".tmp_thumbs"

# ===敏感词管理===
BAN_WORDS_PATH = BASE_DIR / "ban_words.json"

THUMB_SUFFIX = "_thumb"
THUMB_OUTPUT_SIZE = 300
THUMB_QUALITY = 90
ROW_THUMB_SIZE = 48
PREVIEW_SIZE = 320

GENDER_OPTIONS = ["male", "female", "unknown"]
STATUS_OPTIONS = ["normal", "star", "lost", "adopted"]
GENDER_LABEL = {"male": "♂ 男孩", "female": "♀ 女孩", "unknown": "❓ 未知"}
GENDER_COLOR = {"male": COLOR_MALE, "female": COLOR_FEMALE, "unknown": COLOR_UNKNOWN}
STATUS_LABEL = {"normal": "正常", "star": "⭐ 喵星", "lost": "🔍 失踪", "adopted": "🏡 被领养"}
STATUS_BADGE = {"normal": "", "star": "⭐", "lost": "🔍", "adopted": "🏡"}

COLUMN_KEYWORDS = {
    "id": "编号", "name": "姓名", "gender": "性别", "affection": "亲人指数",
    "status": "状态", "desc": "概要", "story": "故事", "pic": "图名",
}

PRESET_COLORS = [
    "#F9A8D4", "#A78BFA", "#6EE7B7", "#FCA5A5",
    "#D4A373", "#7F8C8D", "#A3B18A", "#B5838D",
    "#FFB347", "#B5EAD7", "#C7CEEA", "#FF9AA2",
    "#FF7E93", "#FF9AAD", "#FFB6C6", "#FFD1DC",
    "#5B8DEF", "#E57373", "#81C784", "#FFB74D",
]


# ============================================================
# Excel 读写
# ============================================================
class ExcelStore:
    def __init__(self, path):
        self.path = path
        self.wb = None
        self.ws = None
        self.col_idx = {}

    def load(self):
        if not self.path.exists():
            raise FileNotFoundError(f"找不到 {self.path}")
        self.wb = load_workbook(self.path)
        self.ws = self.wb.active
        headers = {}
        for c in range(1, self.ws.max_column + 1):
            v = self.ws.cell(row=1, column=c).value
            if v is None:
                continue
            headers[c] = str(v).strip()
        for key, kw in COLUMN_KEYWORDS.items():
            matched = [c for c, h in headers.items() if kw in h]
            if not matched:
                raise KeyError(f"Excel 中找不到含 '{kw}' 的列")
            if len(matched) > 1:
                raise ValueError(f"列 '{kw}' 匹配到多个: {matched}")
            self.col_idx[key] = matched[0]

    def all_rows(self):
        rows = []
        for r in range(2, self.ws.max_row + 1):
            raw_id = self.ws.cell(row=r, column=self.col_idx["id"]).value
            if raw_id is None or str(raw_id).strip() == "":
                continue
            rows.append({
                "_row": r,
                "id": self._fmt_id(raw_id),
                "name": self._cell_str(r, "name"),
                "gender": self._cell_str(r, "gender") or "unknown",
                "affection": self._cell_int(r, "affection", default=1),
                "status": self._cell_str(r, "status") or "normal",
                "desc": self._cell_str(r, "desc"),
                "story": self._cell_str(r, "story"),
                "pic_name": self._cell_str(r, "pic"),
            })
        return rows

    def _cell_str(self, r, key):
        v = self.ws.cell(row=r, column=self.col_idx[key]).value
        return "" if v is None else str(v).strip()

    def _cell_int(self, r, key, default=1):
        v = self.ws.cell(row=r, column=self.col_idx[key]).value
        if v is None or str(v).strip() == "":
            return default
        try:
            return int(v)
        except Exception:
            return default

    @staticmethod
    def _fmt_id(raw):
        if isinstance(raw, (int, float)):
            return f"{int(raw):02d}"
        return str(raw).strip().zfill(2)

    def update_row(self, row_index, fields):
        for key, val in fields.items():
            if key in self.col_idx:
                self.ws.cell(row=row_index, column=self.col_idx[key]).value = val

    def find_row_by_id(self, cat_id):
        for r in range(2, self.ws.max_row + 1):
            raw = self.ws.cell(row=r, column=self.col_idx["id"]).value
            if raw is None:
                continue
            if self._fmt_id(raw) == cat_id:
                return r
        return None

    def append_row(self, fields):
        new_r = self.ws.max_row + 1
        while self.ws.cell(row=new_r, column=self.col_idx["id"]).value not in (None, ""):
            new_r += 1
        self.update_row(new_r, fields)
        return new_r

    def delete_row(self, row_index):
        self.ws.delete_rows(row_index, 1)

    def next_id(self):
        max_id = 0
        for r in range(2, self.ws.max_row + 1):
            raw = self.ws.cell(row=r, column=self.col_idx["id"]).value
            if raw is None:
                continue
            try:
                max_id = max(max_id, int(str(raw).strip()))
            except Exception:
                pass
        return f"{max_id + 1:02d}"

    def save(self):
        self.wb.save(self.path)

    def rewrite_rows(self, rows):
        while self.ws.max_row > 1:
            self.ws.delete_rows(2)
        for idx, cat in enumerate(rows, start=2):
            self.ws.cell(row=idx, column=self.col_idx["id"]).value = int(cat["id"])
            self.ws.cell(row=idx, column=self.col_idx["name"]).value = cat["name"]
            self.ws.cell(row=idx, column=self.col_idx["gender"]).value = cat["gender"]
            self.ws.cell(row=idx, column=self.col_idx["affection"]).value = cat["affection"]
            self.ws.cell(row=idx, column=self.col_idx["status"]).value = cat["status"]
            self.ws.cell(row=idx, column=self.col_idx["desc"]).value = cat["desc"]
            self.ws.cell(row=idx, column=self.col_idx["story"]).value = cat["story"]
            self.ws.cell(row=idx, column=self.col_idx["pic"]).value = cat["pic_name"]


# ============================================================
# cats.json 与标签相关工具
# ============================================================
def regenerate_cats_json(rows, classified_dir, json_path, tags_map=None):
    warnings = []
    cats = []
    for row in rows:
        cat_id = row["id"]
        name = row["name"]
        pic_name = row["pic_name"]
        folder_name = f"{cat_id} {name}"
        folder_real = classified_dir / folder_name
        folder_json = f"classified/{folder_name}"

        if not folder_real.is_dir():
            warnings.append(f"[{cat_id}] 文件夹缺失: {folder_name}")
            continue
        if not pic_name:
            warnings.append(f"[{cat_id}] 图名为空")
            continue

        avatar = f"{folder_json}/{pic_name}_01_thumb.jpg"
        avatar_hd = f"{folder_json}/{pic_name}_01.jpg"
        if not (folder_real / f"{pic_name}_01.jpg").is_file():
            warnings.append(f"[{cat_id}] 缺头像原图 {pic_name}_01.jpg")
        if not (folder_real / f"{pic_name}_01_thumb.jpg").is_file():
            warnings.append(f"[{cat_id}] 缺头像缩略图 {pic_name}_01_thumb.jpg")

        other_photos = []
        pattern = str(folder_real / f"{pic_name}_[0-9][0-9].jpg")
        for fpath in sorted(glob.glob(pattern)):
            m = re.match(rf"^{re.escape(pic_name)}_(\d{{2}})\.jpg$", os.path.basename(fpath))
            if not m:
                continue
            seq = int(m.group(1))
            if seq < 2:
                continue
            seq_str = f"{seq:02d}"
            thumb_file = folder_real / f"{pic_name}_{seq_str}_thumb.jpg"
            if not thumb_file.is_file():
                warnings.append(f"[{cat_id}] 缺缩略图 {thumb_file.name}")
            other_photos.append({
                "thumb": f"{folder_json}/{pic_name}_{seq_str}_thumb.jpg",
                "hd": f"{folder_json}/{pic_name}_{seq_str}.jpg",
            })

        tags = []
        if tags_map and cat_id in tags_map:
            tags = [t for t in tags_map[cat_id] if isinstance(t, str) and t.strip()]

        cats.append({
            "id": cat_id, "name": name,
            "gender": row["gender"] or "unknown",
            "avatar": avatar, "avatar_hd": avatar_hd,
            "affection": row["affection"],
            "status": row["status"] or "normal",
            "desc": row["desc"], "story": row["story"],
            "tags": tags,
            "otherPhotos": other_photos,
        })

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(cats, f, ensure_ascii=False, indent=2)
    return len(cats), warnings


def load_tags_from_json(json_path):
    tags_map = {}
    if not json_path.exists():
        return tags_map
    try:
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        for cat in data:
            if "id" in cat and "tags" in cat and isinstance(cat["tags"], list):
                tags_map[cat["id"]] = [t.strip() for t in cat["tags"] if isinstance(t, str) and t.strip()]
    except Exception:
        pass
    return tags_map


def load_tag_colors(path):
    if not path.exists():
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_tag_colors(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ============================================================
# 缩略图工具
# ============================================================
def auto_thumbnail(src_path, dst_path, output_size=THUMB_OUTPUT_SIZE, crop_ratio=0.6):
    img = Image.open(src_path).convert("RGB")
    w, h = img.size
    crop = max(100, int(min(w, h) * crop_ratio))
    left = (w - crop) // 2
    top = (h - crop) // 2
    img = img.crop((left, top, left + crop, top + crop))
    img = img.resize((output_size, output_size), Image.LANCZOS)
    img.save(dst_path, "JPEG", quality=THUMB_QUALITY, optimize=True)


def next_seq_for(folder, pic_name):
    if not folder.is_dir():
        return 1
    used = set()
    for f in folder.iterdir():
        m = re.match(rf"^{re.escape(pic_name)}_(\d{{2}})(?:_thumb)?\.jpe?g$",
                     f.name, re.IGNORECASE)
        if m:
            used.add(int(m.group(1)))
    n = 1
    while n in used:
        n += 1
    return n


def open_in_explorer(path: Path):
    try:
        if sys.platform == "win32":
            os.startfile(str(path))
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception as e:
        messagebox.showerror("打不开", str(e))


def detect_fonts(root):
    global FONT_FAMILY_UI, FONT_FAMILY_MONO
    available = set(tkfont.families(root))
    for name in FONT_CANDIDATES_UI:
        if name in available:
            FONT_FAMILY_UI = name
            break
    for name in FONT_CANDIDATES_MONO:
        if name in available:
            FONT_FAMILY_MONO = name
            break

    default = tkfont.nametofont("TkDefaultFont")
    if FONT_FAMILY_UI:
        default.configure(family=FONT_FAMILY_UI, size=10)
    else:
        default.configure(size=10)
    text_font = tkfont.nametofont("TkTextFont")
    if FONT_FAMILY_UI:
        text_font.configure(family=FONT_FAMILY_UI, size=10)
    fixed = tkfont.nametofont("TkFixedFont")
    if FONT_FAMILY_MONO:
        fixed.configure(family=FONT_FAMILY_MONO, size=10)


def ui_font(size=FS_BODY, weight="normal"):
    if FONT_FAMILY_UI:
        return (FONT_FAMILY_UI, size, weight)
    return ("TkDefaultFont", size, weight)


def mono_font(size=FS_SMALL):
    if FONT_FAMILY_MONO:
        return (FONT_FAMILY_MONO, size)
    return ("TkFixedFont", size)


# ============================================================
# 缩略图裁剪器（保持原有）
# ============================================================
class ThumbEditor(tk.Toplevel):
    def __init__(self, master, original_path, on_done=None, on_close=None):
        super().__init__(master)
        self.title(f"编辑缩略图 - {original_path.name}")
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        self.geometry(f"{int(sw*0.8)}x{int(sh*0.8)}")
        self.configure(bg=COLOR_BG)
        self.transient(master)
        self.grab_set()

        self.original_path = original_path
        self.on_done = on_done
        self.on_close = on_close
        self.original_image = Image.open(original_path)
        self.tk_image = None
        self.crop_box = None
        self.display_crop_box = None
        self.display_scale = 1.0
        self.img_x = 0; self.img_y = 0
        self.dragging = False; self.resizing = False
        self.resize_corner = None
        self.offset_x = 0; self.offset_y = 0
        self.crop_w = 0; self.crop_h = 0
        self.min_crop_size = 100
        self.output_size = THUMB_OUTPUT_SIZE

        self._build_ui()
        self._init_crop_box()
        self.after(80, self._display)
        self.bind("<Configure>", self._on_resize)
        self.protocol("WM_DELETE_WINDOW", self._on_window_close)

    def _on_window_close(self):
        if self.on_close:
            self.on_close()
        self.destroy()

    def _build_ui(self):
        toolbar = tk.Frame(self, bg=COLOR_ACCENT, height=56)
        toolbar.pack(fill=tk.X); toolbar.pack_propagate(False)
        tk.Label(toolbar, text=f"✏️  编辑缩略图：{self.original_path.name}",
                 font=ui_font(FS_BODY+2, "bold"), bg=COLOR_ACCENT, fg="white"
                 ).pack(side=tk.LEFT, padx=18)
        tk.Button(toolbar, text="💾 保存覆盖", command=self._save,
                  font=ui_font(FS_BODY, "bold"), bg=COLOR_WARN, fg="white",
                  padx=18, pady=6, relief=tk.FLAT, cursor="hand2", borderwidth=0
                  ).pack(side=tk.RIGHT, padx=10, pady=10)
        tk.Button(toolbar, text="↺ 重置", command=self._reset,
                  font=ui_font(FS_BODY), bg="#fff8d6", fg=COLOR_TEXT,
                  padx=14, pady=6, relief=tk.FLAT, cursor="hand2", borderwidth=0
                  ).pack(side=tk.RIGHT, padx=4, pady=10)
        tk.Button(toolbar, text="🎯 居中", command=self._center,
                  font=ui_font(FS_BODY), bg="#e3f2fd", fg=COLOR_TEXT,
                  padx=14, pady=6, relief=tk.FLAT, cursor="hand2", borderwidth=0
                  ).pack(side=tk.RIGHT, padx=4, pady=10)

        info_bar = tk.Frame(self, bg=COLOR_DIVIDER, height=32)
        info_bar.pack(fill=tk.X); info_bar.pack_propagate(False)
        self.info_label = tk.Label(info_bar, text="拖动方框移动，拖角调整大小（保持 1:1）",
                                   bg=COLOR_DIVIDER, fg=COLOR_SUBTEXT, font=ui_font(FS_SMALL))
        self.info_label.pack(side=tk.LEFT, padx=14)

        self.canvas = tk.Canvas(self, bg="#f0f0f0", highlightthickness=0)
        self.canvas.pack(fill=tk.BOTH, expand=True)
        self.canvas.bind("<ButtonPress-1>", self._on_down)
        self.canvas.bind("<B1-Motion>", self._on_drag)
        self.canvas.bind("<ButtonRelease-1>", self._on_up)

    def _init_crop_box(self):
        w, h = self.original_image.size
        side = max(100, int(min(w, h) * 0.6))
        x1 = (w - side) // 2; y1 = (h - side) // 2
        self.crop_box = (x1, y1, x1 + side, y1 + side)

    def _display(self):
        cw = self.canvas.winfo_width() or 800
        ch = self.canvas.winfo_height() or 580
        iw, ih = self.original_image.size
        self.display_scale = min(cw / iw, ch / ih, 1.0)
        dw, dh = int(iw * self.display_scale), int(ih * self.display_scale)
        disp = self.original_image.resize((dw, dh), Image.LANCZOS)
        self.tk_image = ImageTk.PhotoImage(disp)
        self.canvas.delete("all")
        self.img_x = (cw - dw) // 2
        self.img_y = (ch - dh) // 2
        self.canvas.create_image(self.img_x, self.img_y, anchor=tk.NW, image=self.tk_image)
        self._draw_box()

    def _draw_box(self):
        if not self.crop_box:
            return
        x1, y1, x2, y2 = self.crop_box
        s = self.display_scale
        dx1 = self.img_x + int(x1 * s); dy1 = self.img_y + int(y1 * s)
        dx2 = self.img_x + int(x2 * s); dy2 = self.img_y + int(y2 * s)
        self.display_crop_box = (dx1, dy1, dx2, dy2)
        cw = self.canvas.winfo_width(); ch = self.canvas.winfo_height()
        self.canvas.delete("box")
        for rect in [(0, 0, dx1, ch), (dx2, 0, cw, ch),
                     (dx1, 0, dx2, dy1), (dx1, dy2, dx2, ch)]:
            self.canvas.create_rectangle(*rect, fill="#000", stipple="gray50",
                                         outline="", tags="box")
        self.canvas.create_rectangle(dx1, dy1, dx2, dy2,
                                     outline=COLOR_INFO, width=3, tags="box")
        for hx, hy in [(dx1, dy1), (dx2, dy1), (dx1, dy2), (dx2, dy2)]:
            self.canvas.create_oval(hx - 7, hy - 7, hx + 7, hy + 7,
                                    fill="#FF7043", outline="white", width=2, tags="box")
        side = x2 - x1
        self.info_label.config(text=f"裁剪 {side}×{side}  ▸  输出 {self.output_size}×{self.output_size}")

    def _on_down(self, e):
        if not self.display_crop_box: return
        dx1, dy1, dx2, dy2 = self.display_crop_box
        for name, (cx, cy) in [("nw", (dx1, dy1)), ("ne", (dx2, dy1)),
                                ("sw", (dx1, dy2)), ("se", (dx2, dy2))]:
            if abs(e.x - cx) < 12 and abs(e.y - cy) < 12:
                self.resizing = True; self.resize_corner = name
                self.offset_x, self.offset_y = e.x, e.y
                return
        if dx1 <= e.x <= dx2 and dy1 <= e.y <= dy2:
            self.dragging = True
            self.offset_x = e.x - dx1; self.offset_y = e.y - dy1
            self.crop_w = self.crop_box[2] - self.crop_box[0]
            self.crop_h = self.crop_box[3] - self.crop_box[1]

    def _on_drag(self, e):
        if self.dragging: self._move(e)
        elif self.resizing: self._resize(e)

    def _on_up(self, _):
        self.dragging = False; self.resizing = False; self.resize_corner = None

    def _move(self, e):
        s = self.display_scale
        new_dx1 = e.x - self.offset_x; new_dy1 = e.y - self.offset_y
        x1f = (new_dx1 - self.img_x) / s; y1f = (new_dy1 - self.img_y) / s
        iw, ih = self.original_image.size
        x1 = max(0, min(int(x1f), iw - self.crop_w))
        y1 = max(0, min(int(y1f), ih - self.crop_h))
        self.crop_box = (x1, y1, x1 + self.crop_w, y1 + self.crop_h)
        self._draw_box()

    def _resize(self, e):
        x1, y1, x2, y2 = self.crop_box
        iw, ih = self.original_image.size
        s = self.display_scale
        dxr = (e.x - self.offset_x) / s; dyr = (e.y - self.offset_y) / s
        if self.resize_corner == "nw": x1 += dxr; y1 += dyr
        elif self.resize_corner == "ne": x2 += dxr; y1 += dyr
        elif self.resize_corner == "sw": x1 += dxr; y2 += dyr
        elif self.resize_corner == "se": x2 += dxr; y2 += dyr
        x1 = max(0, x1); y1 = max(0, y1)
        x2 = min(iw, x2); y2 = min(ih, y2)
        if self.resize_corner == "se": fx, fy = x1, y1
        elif self.resize_corner == "sw": fx, fy = x2, y1
        elif self.resize_corner == "ne": fx, fy = x1, y2
        else: fx, fy = x2, y2
        side = max(min(x2 - x1, y2 - y1), self.min_crop_size / s)
        if self.resize_corner == "se":
            x1, y1 = fx, fy; x2, y2 = x1 + side, y1 + side
        elif self.resize_corner == "sw":
            x2, y1 = fx, fy; x1, y2 = x2 - side, y1 + side
        elif self.resize_corner == "ne":
            x1, y2 = fx, fy; x2, y1 = x1 + side, y2 - side
        else:
            x2, y2 = fx, fy; x1, y1 = x2 - side, y2 - side
        x1 = max(0, x1); y1 = max(0, y1)
        x2 = min(iw, x2); y2 = min(ih, y2)
        self.crop_box = (int(x1), int(y1), int(x2), int(y2))
        self._draw_box()
        self.offset_x, self.offset_y = e.x, e.y

    def _on_resize(self, _e):
        if self.original_image and self.tk_image:
            self._display()

    def _reset(self):
        self._init_crop_box(); self._display()

    def _center(self):
        x1, y1, x2, y2 = self.crop_box
        size = x2 - x1
        iw, ih = self.original_image.size
        nx = (iw - size) // 2; ny = (ih - size) // 2
        self.crop_box = (nx, ny, nx + size, ny + size)
        self._display()

    def _save(self):
        try:
            cropped = self.original_image.crop(self.crop_box)
            thumb = cropped.resize((self.output_size, self.output_size), Image.LANCZOS)
            stem = self.original_path.stem
            ext = self.original_path.suffix
            out_path = self.original_path.with_name(f"{stem}{THUMB_SUFFIX}{ext}")
            thumb.save(out_path, "JPEG", quality=THUMB_QUALITY, optimize=True)
            if self.on_done:
                self.on_done(out_path)
            if self.on_close:
                self.on_close()
            self.destroy()
        except Exception as e:
            messagebox.showerror("失败", f"{e}\n\n{traceback.format_exc()}")


# ============================================================
# 标签综合管理窗口（编辑名称/颜色 + 查看关联猫 + 批量添加标签）
# ============================================================
class TagManageWindow(tk.Toplevel):
    def __init__(self, master, tag_name, tag_color, all_cats, cat_tags_map, all_tag_names,
                 on_save_and_apply):
        """
        on_save_and_apply: 回调函数，接收 (old_name, new_name, new_color, status_dict)
                           status_dict 为 {cat_id: bool}，True 表示拥有标签，False 表示不拥有
                           在此回调中应同时处理标签更新和标签分配。
        """
        super().__init__(master)
        self.tag_name = tag_name
        self.tag_color = tag_color
        self.all_cats = all_cats
        self.cat_tags_map = cat_tags_map
        self.all_tag_names = all_tag_names
        self.on_save_and_apply = on_save_and_apply

        self.title(f"管理标签「{tag_name}」")
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        w = int(sw * 0.7)
        h = int(sh * 0.8)
        x = (sw - w) // 2
        y = (sh - h) // 4
        self.geometry(f"{w}x{h}+{x}+{y}")
        self.minsize(700, 600)
        self.configure(bg=COLOR_BG)
        self.transient(master)
        self.grab_set()

        self._build()

    def _build(self):
        # ── 编辑标签区域 ──
        edit_frame = tk.LabelFrame(self, text="编辑标签", font=ui_font(FS_BODY, "bold"),
                                   bg=COLOR_BG, fg=COLOR_ACCENT)
        edit_frame.pack(fill=tk.X, padx=10, pady=(10, 5))

        row1 = tk.Frame(edit_frame, bg=COLOR_BG)
        row1.pack(fill=tk.X, padx=10, pady=5)
        tk.Label(row1, text="名称：", bg=COLOR_BG, font=ui_font(FS_BODY)).pack(side=tk.LEFT)
        self.name_var = tk.StringVar(value=self.tag_name)
        tk.Entry(row1, textvariable=self.name_var, font=ui_font(FS_BODY), width=20).pack(side=tk.LEFT, padx=5)

        row2 = tk.Frame(edit_frame, bg=COLOR_BG)
        row2.pack(fill=tk.X, padx=10, pady=5)
        tk.Label(row2, text="颜色：", bg=COLOR_BG, font=ui_font(FS_BODY)).pack(side=tk.LEFT)
        self.color_var = tk.StringVar(value=self.tag_color)
        self.color_preview = tk.Label(row2, text="    ", bg=self.tag_color, relief=tk.SOLID,
                                      width=4, borderwidth=2)
        self.color_preview.pack(side=tk.LEFT, padx=5)
        tk.Button(row2, text="🎨 调色盘", command=self._pick_color,
                  font=ui_font(FS_SMALL), bg=COLOR_ACCENT_LIGHT, fg="white",
                  relief=tk.FLAT, padx=8, pady=2, cursor="hand2").pack(side=tk.LEFT, padx=5)

        color_grid = tk.Frame(edit_frame, bg=COLOR_BG)
        color_grid.pack(fill=tk.X, padx=10, pady=5)
        self.color_buttons = []
        for i, c in enumerate(PRESET_COLORS):
            btn = tk.Button(color_grid, bg=c, width=3, height=1, relief=tk.FLAT,
                            borderwidth=1, cursor="hand2",
                            command=lambda col=c: self._select_color(col))
            btn.grid(row=i // 10, column=i % 10, padx=2, pady=2)
            self.color_buttons.append((btn, c))
        self._highlight_selected_color(self.tag_color)

        # ── 猫咪列表 ──
        cats_frame = tk.LabelFrame(self, text="关联猫咪（勾选 = 拥有此标签）", font=ui_font(FS_BODY, "bold"),
                                   bg=COLOR_BG, fg=COLOR_ACCENT)
        cats_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 5))

        list_container = tk.Frame(cats_frame, bg=COLOR_BG)
        list_container.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        canvas = tk.Canvas(list_container, bg=COLOR_BG, highlightthickness=0)
        scrollbar = ttk.Scrollbar(list_container, orient="vertical", command=canvas.yview)
        self.cats_list_frame = tk.Frame(canvas, bg=COLOR_BG)
        self.cats_list_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.cats_list_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self.cat_check_vars = {}
        self._populate_cats_list()

        # ── 底部统一按钮 ──
        bottom_bar = tk.Frame(self, bg=COLOR_BG)
        bottom_bar.pack(fill=tk.X, padx=10, pady=(0, 10))
        tk.Button(bottom_bar, text="💾 保存并应用", command=self._save_and_apply,
                  font=ui_font(FS_BODY, "bold"), bg=COLOR_ACCENT, fg="white",
                  relief=tk.FLAT, padx=20, pady=6, cursor="hand2").pack(side=tk.RIGHT, padx=5)

    def _populate_cats_list(self):
        for w in self.cats_list_frame.winfo_children():
            w.destroy()
        self.cat_check_vars.clear()
        for cat in self.all_cats:
            cat_id = cat["id"]
            cat_name = cat["name"]
            has_tag = self.tag_name in self.cat_tags_map.get(cat_id, [])
            text = f"{cat_id} {cat_name}"
            if has_tag:
                text += " ★"
            var = tk.BooleanVar(value=has_tag)
            cb = tk.Checkbutton(self.cats_list_frame, text=text, variable=var,
                                bg=COLOR_BG, font=ui_font(FS_SMALL), anchor="w",
                                activebackground=COLOR_BG, selectcolor=COLOR_BG)
            cb.pack(fill=tk.X, padx=5, pady=1)
            self.cat_check_vars[cat_id] = var

    def _select_color(self, color):
        self.color_var.set(color)
        self.color_preview.config(bg=color)
        self._highlight_selected_color(color)

    def _highlight_selected_color(self, color):
        for btn, c in self.color_buttons:
            if c == color:
                btn.config(relief=tk.SOLID, borderwidth=2)
            else:
                btn.config(relief=tk.FLAT, borderwidth=1)

    def _pick_color(self):
        chosen = colorchooser.askcolor(color=self.color_var.get(), parent=self)
        if chosen[1]:
            self._select_color(chosen[1])

    def _save_and_apply(self):
        new_name = self.name_var.get().strip()
        if not new_name:
            messagebox.showwarning("提示", "名称不能为空", parent=self)
            return
        # 检查名称是否与其他标签冲突（除了自身）
        if new_name != self.tag_name and new_name in self.all_tag_names:
            messagebox.showwarning("冲突", f"标签「{new_name}」已存在", parent=self)
            return
        new_color = self.color_var.get()
        status = {cat_id: var.get() for cat_id, var in self.cat_check_vars.items()}

        # 调用回调：传递旧名称、新名称、新颜色、状态字典
        self.on_save_and_apply(self.tag_name, new_name, new_color, status)

        # 更新窗口标题和内部缓存的标签名
        self.tag_name = new_name
        self.title(f"管理标签「{new_name}」")
        messagebox.showinfo("成功", "标签信息与分配已更新", parent=self)
        # 刷新猫咪列表（反映新的标签分配和名称）
        self._populate_cats_list()

# ============================================================
# 敏感词管理窗口
# ============================================================
class BanWordsEditor(tk.Toplevel):
    def __init__(self, master, json_path):
        super().__init__(master)
        self.json_path = json_path
        self.title("敏感词管理")
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        w, h = int(sw * 0.55), int(sh * 0.65)      # 宽度取屏幕 55%，高度 65%
        x, y = (sw - w) // 2, (sh - h) // 2        # 居中
        self.geometry(f"{w}x{h}+{x}+{y}")
        self.minsize(600, 500)                     # 最小尺寸防止过小
        self.configure(bg=COLOR_BG)
        self.transient(master)
        self.grab_set()

        self.words = self._load()

        # 标题
        tk.Label(self, text="敏感词列表", font=ui_font(FS_BODY+2, "bold"),
                 bg=COLOR_BG, fg=COLOR_ACCENT).pack(pady=(14, 6))

        # 列表区
        list_frame = tk.Frame(self, bg=COLOR_BG)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=12, pady=4)

        self.listbox = tk.Listbox(list_frame, font=ui_font(FS_BODY), activestyle="dotbox",
                                  selectbackground=COLOR_HOVER, selectforeground=COLOR_TEXT,
                                  relief=tk.FLAT, bg="white", highlightthickness=1,
                                  highlightbackground=COLOR_DIVIDER)
        self.listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.listbox.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.listbox.config(yscrollcommand=scrollbar.set)
        self._refresh_list()

        # 输入与按钮
        ctrl = tk.Frame(self, bg=COLOR_BG)
        ctrl.pack(fill=tk.X, padx=12, pady=10)

        self.entry_var = tk.StringVar()
        tk.Entry(ctrl, textvariable=self.entry_var, font=ui_font(FS_BODY), width=28,
                 relief=tk.FLAT, bg="#fafafa", highlightthickness=1,
                 highlightbackground=COLOR_DIVIDER, highlightcolor=COLOR_INFO).pack(side=tk.LEFT, padx=(0, 6))

        tk.Button(ctrl, text="添加", font=ui_font(FS_SMALL, "bold"), bg=COLOR_OK, fg="white",
                  relief=tk.FLAT, padx=10, pady=5, cursor="hand2", command=self._add).pack(side=tk.LEFT, padx=2)
        tk.Button(ctrl, text="删除选中", font=ui_font(FS_SMALL, "bold"), bg=COLOR_DANGER, fg="white",
                  relief=tk.FLAT, padx=10, pady=5, cursor="hand2", command=self._remove).pack(side=tk.LEFT, padx=2)

        # 底部保存按钮
        tk.Button(self, text="💾 保存", font=ui_font(FS_BODY, "bold"), bg=COLOR_ACCENT, fg="white",
                  relief=tk.FLAT, padx=24, pady=8, cursor="hand2", command=self._save).pack(pady=(4, 16))

    def _load(self):
        if not self.json_path.exists():
            return []
        try:
            with open(self.json_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            return data.get("words", [])
        except Exception:
            return []

    def _refresh_list(self):
        self.listbox.delete(0, tk.END)
        for w in self.words:
            self.listbox.insert(tk.END, w)

    def _add(self):
        new_word = self.entry_var.get().strip()
        if not new_word:
            messagebox.showwarning("提示", "请输入敏感词", parent=self)
            return
        if new_word in self.words:
            messagebox.showwarning("提示", "该词已存在", parent=self)
            return
        self.words.append(new_word)
        self.entry_var.set("")
        self._refresh_list()

    def _remove(self):
        sel = self.listbox.curselection()
        if not sel:
            messagebox.showwarning("提示", "请先选中一个词", parent=self)
            return
        idx = sel[0]
        removed = self.words.pop(idx)
        self._refresh_list()

    def _save(self):
        try:
            with open(self.json_path, 'w', encoding='utf-8') as f:
                json.dump({"words": self.words}, f, ensure_ascii=False, indent=2)
            messagebox.showinfo("成功", "敏感词已保存", parent=self)
            self.destroy()
        except Exception as e:
            messagebox.showerror("保存失败", str(e), parent=self)


# ============================================================
# CatEditor（窗口自适应，增加应用标签按钮）
# ============================================================
class CatEditor(tk.Toplevel):
    def __init__(self, master, store, cat=None, on_saved=None, on_tags_applied=None,
                 tag_colors=None, cat_tags=None):
        super().__init__(master)
        self.store = store
        self.cat = cat
        self.on_saved = on_saved
        self.on_tags_applied = on_tags_applied
        self.is_new = cat is None
        self.tag_colors = tag_colors or []
        self.cat_tags = cat_tags or []

        title = "新增猫咪" if self.is_new else f"编辑 {cat['id']} {cat['name']}"
        self.title(title)

        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        # 窗口宽度取屏幕宽度的 85%，高度取 80%，最小尺寸也相应调整
        target_w = max(1100, int(sw * 0.85))
        target_h = max(750, int(sh * 0.80))
        self.geometry(f"{target_w}x{target_h}")
        self.minsize(max(900, int(sw*0.55)), max(700, int(sh*0.55)))

        self.configure(bg=COLOR_BG)
        self.transient(master)
        self.grab_set()
        self.protocol("WM_DELETE_WINDOW", self._on_close_request)

        self.new_photo_queue = []
        self._initial_snapshot = None
        self._preview_original_img = None
        self._preview_thumb_img = None
        self._last_original = None
        self._last_thumb = None
        self.tag_vars = {}

        TEMP_THUMB_DIR.mkdir(parents=True, exist_ok=True)

        self._build()
        self._load_values()
        if not self.is_new:
            self._refresh_existing_photos()
        self._initial_snapshot = self._snapshot()
        self._made_fs_changes = False

    def _snapshot(self):
        return (
            self.var_id.get(), self.var_name.get(), self.var_pic.get(),
            self.var_gender.get(), self.var_status.get(),
            int(self.var_affection.get() or 0), self.var_desc.get(),
            self.txt_story.get("1.0", tk.END),
            tuple((str(p["src"]), p["seq"], p.get("thumb_temp", None) and str(p["thumb_temp"])) for p in self.new_photo_queue),
            {t: v.get() for t, v in self.tag_vars.items()}
        )

    def _is_dirty(self):
        if self._initial_snapshot is None:
            return False
        return self._snapshot() != self._initial_snapshot or self._made_fs_changes

    def _mark_dirty(self):
        self._made_fs_changes = True

    def _on_close_request(self):
        if not self._is_dirty():
            self.destroy(); return
        ans = messagebox.askyesnocancel(
            "未保存的修改",
            "你有未保存的修改。\n\n是 ＝ 保存并退出\n否 ＝ 丢弃并退出\n取消 ＝ 继续编辑"
        )
        if ans is None:
            return
        if ans:
            self._save()
        else:
            self.destroy()

    def _build(self):
        head = tk.Frame(self, bg=COLOR_ACCENT, height=58)
        head.pack(fill=tk.X); head.pack_propagate(False)
        title = "➕ 新增一只猫咪" if self.is_new else f"✏️ {self.cat['id']} {self.cat['name']}"
        tk.Label(head, text=title, font=ui_font(FS_BODY+4, "bold"),
                 bg=COLOR_ACCENT, fg="white").pack(side=tk.LEFT, padx=22, pady=14)

        outer = tk.Frame(self, bg=COLOR_BG)
        outer.pack(fill=tk.BOTH, expand=True, padx=20, pady=14)

        left = tk.Frame(outer, bg=COLOR_BG)
        left.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        right = tk.Frame(outer, bg=COLOR_BG)
        right.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(14, 0))

        # ----- 基本信息 -----
        fields = self._make_card(left, "基本信息")
        fields.pack(fill=tk.X, pady=(0, 12))

        self.var_id = tk.StringVar()
        self.var_name = tk.StringVar()
        self.var_pic = tk.StringVar()
        self.var_gender = tk.StringVar(value="unknown")
        self.var_status = tk.StringVar(value="normal")
        self.var_affection = tk.IntVar(value=3)
        self.var_desc = tk.StringVar()

        body = tk.Frame(fields, bg=COLOR_CARD)
        body.pack(fill=tk.X, padx=18, pady=12)
        body.columnconfigure(1, weight=1); body.columnconfigure(3, weight=1)

        def label(parent, text, r, c):
            tk.Label(parent, text=text, bg=COLOR_CARD, fg=COLOR_SUBTEXT,
                     font=ui_font(FS_SMALL)).grid(row=r, column=c, sticky="w", padx=4, pady=6)

        def entry(parent, var, r, c, **kw):
            e = tk.Entry(parent, textvariable=var, font=ui_font(FS_BODY), relief=tk.FLAT,
                         bg="#fafafa", highlightthickness=1, highlightbackground=COLOR_DIVIDER,
                         highlightcolor=COLOR_INFO, **kw)
            e.grid(row=r, column=c, sticky="we", padx=4, pady=6, ipady=4)
            return e

        label(body, "编号", 0, 0)
        e_id = entry(body, self.var_id, 0, 1, width=10)
        if not self.is_new:
            e_id.config(state="disabled")

        label(body, "姓名", 0, 2)
        entry(body, self.var_name, 0, 3)

        label(body, "图名", 1, 0)
        entry(body, self.var_pic, 1, 1)

        label(body, "概要", 1, 2)
        entry(body, self.var_desc, 1, 3)

        label(body, "性别", 2, 0)
        cb_g = ttk.Combobox(body, textvariable=self.var_gender, values=GENDER_OPTIONS, state="readonly", font=ui_font(FS_BODY))
        cb_g.grid(row=2, column=1, sticky="we", padx=4, pady=6, ipady=2)

        label(body, "状态", 2, 2)
        cb_s = ttk.Combobox(body, textvariable=self.var_status, values=STATUS_OPTIONS, state="readonly", font=ui_font(FS_BODY))
        cb_s.grid(row=2, column=3, sticky="we", padx=4, pady=6, ipady=2)

        label(body, "亲人指数", 3, 0)
        aff = tk.Frame(body, bg=COLOR_CARD)
        aff.grid(row=3, column=1, columnspan=3, sticky="w", padx=4, pady=6)
        tk.Spinbox(aff, from_=1, to=5, textvariable=self.var_affection, font=ui_font(FS_BODY), width=5, relief=tk.FLAT,
                   bg="#fafafa", highlightthickness=1, highlightbackground=COLOR_DIVIDER).pack(side=tk.LEFT, ipady=2)
        self.aff_preview = tk.Label(aff, bg=COLOR_CARD, fg=COLOR_PAW, font=ui_font(FS_BODY+2))
        self.aff_preview.pack(side=tk.LEFT, padx=10)
        self.var_affection.trace_add("write", lambda *_: self._update_paw_preview())
        self._update_paw_preview()

        if self.is_new:
            tk.Label(body, text="提示：编号与图名保存后不可修改。图名只能包含 英文/数字/下划线。",
                     bg=COLOR_CARD, fg=COLOR_SUBTEXT, font=ui_font(FS_TINY),
                     wraplength=520, justify="left").grid(row=4, column=0, columnspan=4, sticky="w", padx=4, pady=(8, 0))
        else:
            tk.Label(body, text="提示：图名修改后将重命名所有相关图片文件，请谨慎操作。只能包含英文/数字/下划线。",
                     bg=COLOR_CARD, fg=COLOR_SUBTEXT, font=ui_font(FS_TINY),
                     wraplength=520, justify="left").grid(row=4, column=0, columnspan=4, sticky="w", padx=4, pady=(8, 0))

        # ----- 故事 -----
        story_card = self._make_card(left, "故事")
        story_card.pack(fill=tk.BOTH, expand=True, pady=(0, 0))
        story_inner = tk.Frame(story_card, bg=COLOR_CARD)
        story_inner.pack(fill=tk.BOTH, expand=True, padx=18, pady=12)
        self.txt_story = tk.Text(story_inner, height=8, font=ui_font(FS_BODY), wrap="word", relief=tk.FLAT,
                                 bg="#fafafa", padx=10, pady=10, highlightthickness=1,
                                 highlightbackground=COLOR_DIVIDER, highlightcolor=COLOR_INFO)
        self.txt_story.pack(fill=tk.BOTH, expand=True)

        # ----- 标签选择区（带应用按钮） -----
        tag_card = self._make_card(left, "标签")
        tag_card.pack(fill=tk.X, pady=(8, 0))
        tag_inner = tk.Frame(tag_card, bg=COLOR_CARD)
        tag_inner.pack(fill=tk.BOTH, expand=True, padx=18, pady=12)
        self.tag_frame = tk.Frame(tag_inner, bg=COLOR_CARD)
        self.tag_frame.pack(fill=tk.BOTH)
        self._populate_tag_checkboxes()
        btn_apply_tags = tk.Button(tag_inner, text="✅ 应用标签", command=self._apply_tags_only,
                                   font=ui_font(FS_SMALL, "bold"), bg=COLOR_INFO, fg="white",
                                   relief=tk.FLAT, padx=10, pady=5, cursor="hand2")
        btn_apply_tags.pack(anchor="e", pady=(6, 0))

        # ----- 图片区 -----
        photo_card = self._make_card(right, "图片")
        photo_card.pack(fill=tk.BOTH, expand=True)
        photo_inner = tk.Frame(photo_card, bg=COLOR_CARD)
        photo_inner.pack(fill=tk.BOTH, expand=True, padx=14, pady=12)

        preview_frame = tk.Frame(photo_inner, bg=COLOR_CARD)
        preview_frame.pack(fill=tk.X, pady=(0, 8))
        preview_frame.columnconfigure(0, weight=1)
        preview_frame.columnconfigure(1, weight=1)

        def make_preview_box(parent, column, label_text):
            box = tk.Frame(parent, bg="#fafafa", highlightthickness=1, highlightbackground=COLOR_DIVIDER, width=240, height=240)
            box.grid(row=0, column=column, padx=2, pady=4, sticky="n")
            box.pack_propagate(False)
            lbl = tk.Label(box, bg="#fafafa", fg=COLOR_SUBTEXT, text=label_text, font=ui_font(FS_SMALL))
            lbl.pack(fill=tk.BOTH, expand=True)
            return box, lbl

        self.preview_original_box, self.preview_original_label = make_preview_box(preview_frame, 0, "原图预览")
        self.preview_thumb_box, self.preview_thumb_label = make_preview_box(preview_frame, 1, "缩略图预览")

        def pill(parent, text, color, cmd):
            return tk.Button(parent, text=text, command=cmd, font=ui_font(FS_SMALL, "bold"),
                             bg=color, fg="white", relief=tk.FLAT, borderwidth=0, padx=10, pady=5, cursor="hand2")

        if self.is_new:
            row1 = tk.Frame(photo_inner, bg=COLOR_CARD)
            row1.pack(fill=tk.X, pady=(2, 2))
            pill(row1, "📷 选主图", COLOR_OK, lambda: self._start_photo_selection(seq=1)).pack(side=tk.LEFT, padx=2)
            pill(row1, "➕ 加补充", COLOR_INFO, lambda: self._start_photo_selection(seq=None)).pack(side=tk.LEFT, padx=2)
            row2 = tk.Frame(photo_inner, bg=COLOR_CARD)
            row2.pack(fill=tk.X, pady=(0, 4))
            pill(row2, "✏️ 重裁缩略", COLOR_WARN, self._manual_rethumb_queued).pack(side=tk.LEFT, padx=2)
            pill(row2, "🗑 移除", COLOR_DANGER, self._remove_queued).pack(side=tk.LEFT, padx=2)
        else:
            row1 = tk.Frame(photo_inner, bg=COLOR_CARD)
            row1.pack(fill=tk.X, pady=(2, 2))
            pill(row1, "✏️ 重裁缩略", COLOR_WARN, self._edit_existing_thumb).pack(side=tk.LEFT, padx=2)
            pill(row1, "➕ 补充图", COLOR_INFO, self._add_extra_existing).pack(side=tk.LEFT, padx=2)
            pill(row1, "🗑 删除", COLOR_DANGER, self._delete_existing_photo).pack(side=tk.LEFT, padx=2)
            row2 = tk.Frame(photo_inner, bg=COLOR_CARD)
            row2.pack(fill=tk.X, pady=(0, 4))
            pill(row2, "🔄 设为主图", "#7E57C2", self._swap_to_main).pack(side=tk.LEFT, padx=2)
            pill(row2, "📁 打开文件夹", COLOR_ACCENT_LIGHT, lambda: open_in_explorer(self._folder())).pack(side=tk.LEFT, padx=2)

        list_holder = tk.Frame(photo_inner, bg="white", highlightthickness=1, highlightbackground=COLOR_DIVIDER)
        list_holder.pack(fill=tk.BOTH, expand=True, pady=(4, 0))
        self.photo_list = tk.Listbox(list_holder, font=mono_font(FS_SMALL), activestyle="dotbox", relief=tk.FLAT,
                                     bg="white", borderwidth=0, highlightthickness=0, selectbackground=COLOR_HOVER,
                                     selectforeground=COLOR_TEXT)
        self.photo_list.pack(fill=tk.BOTH, expand=True, side=tk.LEFT)
        sb = ttk.Scrollbar(list_holder, command=self.photo_list.yview)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.photo_list.config(yscrollcommand=sb.set)
        self.photo_list.bind("<<ListboxSelect>>", lambda _e: self._on_photo_select())

        bottom = tk.Frame(self, bg=COLOR_BG)
        bottom.pack(fill=tk.X, padx=20, pady=(0, 16))
        tk.Button(bottom, text="取消", command=self._on_close_request, font=ui_font(FS_BODY), bg="#eaeaea", fg=COLOR_TEXT,
                  relief=tk.FLAT, borderwidth=0, padx=22, pady=8, cursor="hand2").pack(side=tk.RIGHT, padx=4)
        tk.Button(bottom, text="💾 保存并同步", command=self._save, font=ui_font(FS_BODY, "bold"), bg=COLOR_ACCENT, fg="white",
                  relief=tk.FLAT, borderwidth=0, padx=24, pady=8, cursor="hand2").pack(side=tk.RIGHT, padx=4)

    def _apply_tags_only(self):
        if self.is_new:
            messagebox.showwarning("提示", "请先保存猫咪基本信息", parent=self)
            return
        selected = [t for t, var in self.tag_vars.items() if var.get()]
        cat_id = self.cat["id"]
        if self.on_tags_applied:
            self.on_tags_applied(cat_id, selected)
        messagebox.showinfo("已应用", "标签已更新并同步到 cats.json", parent=self)
        self._initial_snapshot = self._snapshot()

    def _populate_tag_checkboxes(self):
        for widget in self.tag_frame.winfo_children():
            widget.destroy()
        self.tag_vars.clear()
        for tag in self.tag_colors:
            var = tk.BooleanVar(value=(tag in self.cat_tags))
            cb = tk.Checkbutton(self.tag_frame, text=tag, variable=var,
                                bg=COLOR_CARD, font=ui_font(FS_SMALL),
                                activebackground=COLOR_CARD, selectcolor=COLOR_CARD)
            cb.pack(side=tk.LEFT, padx=4, pady=2)
            self.tag_vars[tag] = var

    def _make_card(self, parent, title):
        card = tk.Frame(parent, bg=COLOR_CARD, highlightthickness=1, highlightbackground=COLOR_CARD_BORDER)
        title_bar = tk.Frame(card, bg=COLOR_CARD)
        title_bar.pack(fill=tk.X, padx=18, pady=(12, 4))
        tk.Label(title_bar, text=title, bg=COLOR_CARD, fg=COLOR_ACCENT, font=ui_font(FS_BODY+1, "bold")).pack(side=tk.LEFT)
        sep = tk.Frame(card, bg=COLOR_DIVIDER, height=1)
        sep.pack(fill=tk.X, padx=18)
        return card

    def _update_paw_preview(self):
        try: n = int(self.var_affection.get())
        except: n = 0
        n = max(0, min(5, n))
        self.aff_preview.config(text="🐾" * n + "·" * (5 - n))

    def _load_values(self):
        if self.is_new:
            self.var_id.set(self.store.next_id())
            return
        c = self.cat
        self.var_id.set(c["id"])
        self.var_name.set(c["name"])
        self.var_pic.set(c["pic_name"])
        self.var_gender.set(c["gender"] or "unknown")
        self.var_status.set(c["status"] or "normal")
        self.var_affection.set(c["affection"] or 1)
        self.var_desc.set(c["desc"])
        self.txt_story.delete("1.0", tk.END)
        self.txt_story.insert("1.0", c["story"])

    def _show_preview_both(self, original_path, thumb_path=None):
        if original_path and Path(original_path).is_file():
            try:
                img = Image.open(original_path).convert("RGB")
                img.thumbnail((220, 220), Image.LANCZOS)
                self._preview_original_img = ImageTk.PhotoImage(img)
                self.preview_original_label.config(image=self._preview_original_img, text="")
            except Exception:
                self._preview_original_img = None
                self.preview_original_label.config(image="", text="(无法加载)")
        else:
            self._preview_original_img = None
            self.preview_original_label.config(image="", text="(无原图)")

        if thumb_path and Path(thumb_path).is_file():
            try:
                img = Image.open(thumb_path).convert("RGB")
                img.thumbnail((220, 220), Image.LANCZOS)
                self._preview_thumb_img = ImageTk.PhotoImage(img)
                self.preview_thumb_label.config(image=self._preview_thumb_img, text="")
            except Exception:
                self._preview_thumb_img = None
                self.preview_thumb_label.config(image="", text="(无法加载)")
        else:
            self._preview_thumb_img = None
            self.preview_thumb_label.config(image="", text="选中图片后显示")
        self._last_original = original_path
        self._last_thumb = thumb_path

    def _on_photo_select(self):
        idx = self._selected_idx()
        if idx is None: return
        if self.is_new:
            if 0 <= idx < len(self.new_photo_queue):
                item = self.new_photo_queue[idx]
                src = item["src"]
                thumb = item.get("thumb_temp", None)
                self._show_preview_both(str(src), str(thumb) if thumb else None)
        else:
            seq = self._selected_existing_seq()
            if seq is None: return
            folder = self._folder()
            pic = self.cat["pic_name"]
            original = folder / f"{pic}_{seq:02d}.jpg"
            thumb = folder / f"{pic}_{seq:02d}_thumb.jpg" if seq else None
            self._show_preview_both(
                str(original) if original.is_file() else None,
                str(thumb) if (thumb and thumb.is_file()) else None
            )

    def _selected_idx(self):
        sel = self.photo_list.curselection()
        return sel[0] if sel else None

    def _start_photo_selection(self, seq):
        files = filedialog.askopenfilenames(title="选择图片", filetypes=[("图片", "*.jpg *.jpeg *.png")])
        if not files: return
        if seq == 1:
            self.new_photo_queue = [p for p in self.new_photo_queue if p["seq"] != 1]
        self._pending_files = list(files)
        self._pending_seq = seq
        self._process_next_file()

    def _process_next_file(self):
        if not self._pending_files:
            self._refresh_queue_list()
            return
        file_path = self._pending_files.pop(0)
        seq = self._pending_seq
        def on_done(thumb_path):
            tmp_name = f"tmp_{os.path.basename(file_path)}_{os.urandom(4).hex()}.jpg"
            dest = TEMP_THUMB_DIR / tmp_name
            shutil.copyfile(thumb_path, dest)
            self.new_photo_queue.append({"src": Path(file_path), "seq": seq, "thumb_temp": dest})
        def on_close_full():
            if not any(p["src"] == Path(file_path) for p in self.new_photo_queue):
                self.new_photo_queue.append({"src": Path(file_path), "seq": seq, "thumb_temp": None})
            self.after(50, self._process_next_file)
        tmp_original = TEMP_THUMB_DIR / ("orig_" + os.path.basename(file_path))
        shutil.copyfile(file_path, tmp_original)
        editor = ThumbEditor(self, tmp_original, on_done=lambda p: on_done(p), on_close=on_close_full)
        editor.on_close = on_close_full

    def _refresh_queue_list(self):
        self.photo_list.delete(0, tk.END)
        for p in self.new_photo_queue:
            tag = "主图" if p["seq"] == 1 else "补充"
            status = "  [已裁剪]" if p.get("thumb_temp") else ""
            self.photo_list.insert(tk.END, f"[{tag}] {p['src'].name}{status}")

    def _manual_rethumb_queued(self):
        idx = self._selected_idx()
        if idx is None: messagebox.showwarning("提示", "请先选中一张图片"); return
        item = self.new_photo_queue[idx]
        tmp_original = TEMP_THUMB_DIR / ("orig_" + item["src"].name)
        shutil.copyfile(str(item["src"]), tmp_original)
        def on_done(thumb_path):
            tmp_name = f"tmp_{item['src'].name}_{os.urandom(4).hex()}.jpg"
            dest = TEMP_THUMB_DIR / tmp_name
            shutil.copyfile(thumb_path, dest)
            item["thumb_temp"] = dest
            self._refresh_queue_list()
        ThumbEditor(self, tmp_original, on_done=on_done)

    def _remove_queued(self):
        idx = self._selected_idx()
        if idx is None: return
        del self.new_photo_queue[idx]
        self._refresh_queue_list()

    def _folder(self):
        return CLASSIFIED_DIR / f"{self.cat['id']} {self.cat['name']}"

    def _refresh_existing_photos(self):
        self.photo_list.delete(0, tk.END)
        folder = self._folder()
        if not folder.is_dir(): self.photo_list.insert(tk.END, "(文件夹不存在)"); return
        pic = self.cat["pic_name"]
        seqs = []
        for f in folder.iterdir():
            m = re.match(rf"^{re.escape(pic)}_(\d{{2}})\.jpg$", f.name, re.IGNORECASE)
            if m: seqs.append(int(m.group(1)))
        for seq in sorted(set(seqs)):
            tag = "主图" if seq == 1 else "补充"
            thumb_ok = (folder / f"{pic}_{seq:02d}_thumb.jpg").is_file()
            mark = "" if thumb_ok else "  [缺缩略图]"
            self.photo_list.insert(tk.END, f"[{tag}] {pic}_{seq:02d}.jpg{mark}")
        if self.photo_list.size() > 0:
            self.photo_list.selection_set(0)
            self._on_photo_select()

    def _selected_existing_seq(self):
        idx = self._selected_idx()
        if idx is None: return None
        line = self.photo_list.get(idx)
        m = re.search(r"_(\d{2})\.jpg", line)
        return int(m.group(1)) if m else None

    def _edit_existing_thumb(self):
        seq = self._selected_existing_seq()
        if seq is None: messagebox.showwarning("提示", "请先选中一张原图"); return
        folder = self._folder()
        original = folder / f"{self.cat['pic_name']}_{seq:02d}.jpg"
        if not original.is_file(): messagebox.showerror("失败", f"原图不存在: {original.name}"); return
        ThumbEditor(self, original, on_done=lambda _p: (self._refresh_existing_photos(), self._mark_dirty()))

    def _add_extra_existing(self):
        files = filedialog.askopenfilenames(title="选择补充原图", filetypes=[("图片", "*.jpg *.jpeg *.png")])
        if not files: return
        folder = self._folder()
        folder.mkdir(parents=True, exist_ok=True)
        pic = self.cat["pic_name"]
        for src in files:
            seq = next_seq_for(folder, pic)
            if seq < 2:
                seq = 2
                while (folder / f"{pic}_{seq:02d}.jpg").exists():
                    seq += 1
            dst = folder / f"{pic}_{seq:02d}.jpg"
            try:
                Image.open(src).convert("RGB").save(dst, "JPEG", quality=95)
                thumb_dst = folder / f"{pic}_{seq:02d}_thumb.jpg"
                auto_thumbnail(dst, thumb_dst)
            except Exception as e:
                messagebox.showerror("复制失败", f"{src}: {e}")
        self._refresh_existing_photos()
        self._mark_dirty()

    def _delete_existing_photo(self):
        seq = self._selected_existing_seq()
        if seq is None: messagebox.showwarning("提示", "请先选中一张"); return
        if seq == 1: messagebox.showwarning("拒绝", "主图 (_01) 不能在这里删除"); return
        if not messagebox.askyesno("确认", f"删除 _{seq:02d} 原图与缩略图？"): return
        folder = self._folder()
        pic = self.cat["pic_name"]
        for t in [folder / f"{pic}_{seq:02d}.jpg", folder / f"{pic}_{seq:02d}_thumb.jpg"]:
            if t.is_file():
                try: t.unlink()
                except Exception as e: messagebox.showerror("删除失败", f"{t}: {e}")
        self._refresh_existing_photos()
        self._mark_dirty()

    def _swap_to_main(self):
        seq = self._selected_existing_seq()
        if seq is None or seq == 1: return
        folder = self._folder()
        pic = self.cat["pic_name"]
        if not messagebox.askyesno("确认切换", f"将 {pic}_{seq:02d} 设为主图？"): return
        def swap_pair(a_seq, b_seq):
            for suffix in [".jpg", "_thumb.jpg"]:
                a_file = folder / f"{pic}_{a_seq:02d}{suffix}"
                b_file = folder / f"{pic}_{b_seq:02d}{suffix}"
                tmp_file = folder / f"{pic}_{a_seq:02d}{suffix}.tmpswap"
                if a_file.is_file(): a_file.rename(tmp_file)
                if b_file.is_file(): b_file.rename(a_file)
                if tmp_file.is_file(): tmp_file.rename(b_file)
        swap_pair(1, seq)
        self._refresh_existing_photos()
        self._mark_dirty()

    def _save(self):
        try:
            cat_id = self.var_id.get().strip().zfill(2)
            name = self.var_name.get().strip()
            pic = self.var_pic.get().strip()
            gender = (self.var_gender.get().strip() or "unknown")
            status = (self.var_status.get().strip() or "normal")
            try: affection = int(self.var_affection.get())
            except: affection = 1
            desc = self.var_desc.get().strip()
            story = self.txt_story.get("1.0", tk.END).strip()
            selected_tags = [t for t, var in self.tag_vars.items() if var.get()]

            if not name: messagebox.showwarning("提示", "姓名不能为空"); return
            if not pic: messagebox.showwarning("提示", "图名不能为空"); return
            if not re.match(r"^[A-Za-z0-9_]+$", pic):
                messagebox.showwarning("提示", "图名只能含 英文/数字/下划线"); return

            actions = []
            folder = CLASSIFIED_DIR / f"{cat_id} {name}"

            if self.is_new:
                if folder.exists(): messagebox.showerror("失败", f"文件夹已存在：{folder.name}"); return
                if not any(p["seq"] == 1 for p in self.new_photo_queue):
                    messagebox.showwarning("提示", "请先选主图"); return
                folder.mkdir(parents=True, exist_ok=True)
                main_item = next(p for p in self.new_photo_queue if p["seq"] == 1)
                others = [p for p in self.new_photo_queue if p["seq"] != 1]
                self._copy_one(main_item, folder, pic, 1, actions)
                seq = 2
                for it in others:
                    while (folder / f"{pic}_{seq:02d}.jpg").exists(): seq += 1
                    self._copy_one(it, folder, pic, seq, actions)
                    seq += 1
                self.store.append_row({
                    "id": int(cat_id), "name": name, "gender": gender, "affection": affection,
                    "status": status, "desc": desc, "story": story, "pic": pic,
                })
                actions.append(f"追加 Excel 行 编号={cat_id}")
            else:
                old_pic = self.cat["pic_name"]
                if pic != old_pic:
                    old_folder = CLASSIFIED_DIR / f"{self.cat['id']} {self.cat['name']}"
                    if not old_folder.is_dir():
                        messagebox.showerror("失败", f"猫咪文件夹不存在: {old_folder.name}"); return
                    rename_ops = []
                    try:
                        for f in old_folder.iterdir():
                            m = re.match(rf"^{re.escape(old_pic)}_(\d{{2}})(?:_thumb)?\.jpg$", f.name, re.IGNORECASE)
                            if m:
                                seq = m.group(1)
                                is_thumb = "_thumb" in f.name
                                new_name = f"{pic}_{seq}{'_thumb' if is_thumb else ''}.jpg"
                                rename_ops.append((f, old_folder / new_name))
                        for old_path, new_path in rename_ops: old_path.rename(new_path)
                    except Exception as e:
                        for old_path, new_path in reversed(rename_ops):
                            if new_path.exists():
                                try: new_path.rename(old_path)
                                except: pass
                        messagebox.showerror("图名修改失败", f"重命名图片时出错: {e}"); return
                old_folder = self._folder()
                if name != self.cat["name"] and old_folder.is_dir():
                    if folder.exists(): messagebox.showerror("失败", f"目标文件夹已存在：{folder.name}"); return
                    old_folder.rename(folder)
                row_idx = self.store.find_row_by_id(cat_id)
                if row_idx is None: messagebox.showerror("失败", f"Excel 中找不到编号 {cat_id}"); return
                self.store.update_row(row_idx, {
                    "name": name, "gender": gender, "affection": affection,
                    "status": status, "desc": desc, "story": story, "pic": pic,
                })
                actions.append(f"更新 Excel 行 编号={cat_id}")

            self.result_tags = selected_tags
            self.result_id = cat_id
            self._initial_snapshot = self._snapshot()
            if self.on_saved:
                self.on_saved(actions, cat_id, selected_tags)
            self.destroy()
        except Exception as e:
            messagebox.showerror("保存失败", f"{e}\n\n{traceback.format_exc()}")

    def _copy_one(self, item, folder, pic, seq, actions):
        src = item["src"]
        dst = folder / f"{pic}_{seq:02d}.jpg"
        Image.open(src).convert("RGB").save(dst, "JPEG", quality=95)
        actions.append(f"复制原图 → {dst.name}")
        thumb_dst = folder / f"{pic}_{seq:02d}_thumb.jpg"
        if item.get("thumb_temp") and item["thumb_temp"].exists():
            shutil.copyfile(str(item["thumb_temp"]), thumb_dst)
            actions.append(f"使用手动裁剪缩略图 → {thumb_dst.name}")
        else:
            auto_thumbnail(dst, thumb_dst)
            actions.append(f"自动生成缩略图 → {thumb_dst.name}")


# ============================================================
# 主窗口
# ============================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("🐱 SUAT-cats 管理器")
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        self.geometry(f"{int(sw*0.9)}x{int(sh*0.85)}")
        self.minsize(1100, 700)
        self.configure(bg=COLOR_BG)

        detect_fonts(self)
        self._configure_ttk_styles()

        self.store = ExcelStore(EXCEL_PATH)
        self.store.load()
        self.rows = []
        self.thumb_cache = {}
        self.current_filter = "all"
        self.search_term = ""

        self.tag_colors = {}
        self.cat_tags_map = {}
        self._order_dirty = False

        self.http_server = None
        self.http_server_thread = None

        self._build_menu()
        self._build()
        self._load_all_data()

        self.protocol("WM_DELETE_WINDOW", self._on_close_app)

    def _load_all_data(self):
        try:
            self.store.load()
            self.rows = self.store.all_rows()
        except Exception as e:
            messagebox.showerror("加载 Excel 失败", str(e))
            self.rows = []
        self.tag_colors = load_tag_colors(TAG_COLOR_PATH)
        self.cat_tags_map = load_tags_from_json(JSON_PATH)
        self._render()
        self._render_tag_panel()
        self.status_var.set(f"已加载 {len(self.rows)} 只猫咪")
        self._order_dirty = False

    def _on_close_app(self):
        if self._order_dirty:
            ans = messagebox.askyesnocancel(
                "未保存的顺序变更",
                "你移动过猫咪的顺序，但尚未点击「应用变更」保存到 Excel。\n\n"
                "是 ＝ 保存并退出\n否 ＝ 丢弃并退出\n取消 ＝ 继续编辑"
            )
            if ans is None: return
            if ans:
                try: self._apply_changes_silent()
                except Exception as e: messagebox.showerror("保存失败", str(e)); return
        if self.http_server: self._stop_server()
        self.destroy()

    def _apply_changes_silent(self):
        if not self._order_dirty: return
        self._do_apply_changes(silent=True)

    def _do_apply_changes(self, silent=False):
        if not self.rows: return
        old_to_new = {}
        for idx, cat in enumerate(self.rows, start=1):
            new_id = f"{idx:02d}"
            old_id = cat["id"]
            old_to_new[old_id] = new_id

        conflicts = []
        for idx, cat in enumerate(self.rows):
            new_id = f"{idx+1:02d}"
            new_folder_name = f"{new_id} {cat['name']}"
            new_folder = CLASSIFIED_DIR / new_folder_name
            old_folder = CLASSIFIED_DIR / f"{cat['id']} {cat['name']}"
            if new_folder.exists() and new_folder != old_folder:
                conflicts.append(new_folder_name)
        if conflicts:
            messagebox.showerror("重命名冲突", "\n".join(conflicts)); return

        if not silent and not messagebox.askyesno("确认重新编号", "是否按当前顺序重新编号所有猫咪？"):
            return

        temp_suffix = "_renaming"
        for cat in self.rows:
            old_folder = CLASSIFIED_DIR / f"{cat['id']} {cat['name']}"
            if old_folder.is_dir():
                old_folder.rename(CLASSIFIED_DIR / f"{cat['id']} {cat['name']}{temp_suffix}")
        for idx, cat in enumerate(self.rows):
            new_id = f"{idx+1:02d}"
            temp_folder = CLASSIFIED_DIR / f"{cat['id']} {cat['name']}{temp_suffix}"
            if temp_folder.is_dir():
                temp_folder.rename(CLASSIFIED_DIR / f"{new_id} {cat['name']}")

        for idx, cat in enumerate(self.rows):
            cat["id"] = f"{idx+1:02d}"
        self.store.rewrite_rows(self.rows)
        self.store.save()

        new_tags_map = {}
        for old_id, new_id in old_to_new.items():
            if old_id in self.cat_tags_map:
                new_tags_map[new_id] = self.cat_tags_map[old_id]
        self.cat_tags_map = new_tags_map

        count, warns = regenerate_cats_json(self.rows, CLASSIFIED_DIR, JSON_PATH, self.cat_tags_map)
        self._order_dirty = False
        if not silent:
            messagebox.showinfo("应用完成", f"重新编号成功，共 {count} 只猫咪。")
        else:
            self.status_var.set("已自动保存顺序")
        self._reload_from_excel()

    def _apply_changes(self):
        self._do_apply_changes(silent=False)

    # ---------- 标签管理面板 ----------
    def _render_tag_panel(self):
        if not hasattr(self, 'tag_panel_frame'): return
        for w in self.tag_panel_frame.winfo_children():
            w.destroy()

        tk.Label(self.tag_panel_frame, text="🏷️ 标签管理",
                 font=ui_font(FS_BODY+1, "bold"), bg=COLOR_CARD, fg=COLOR_ACCENT
                 ).pack(anchor="w", padx=12, pady=(10, 6))

        list_container = tk.Frame(self.tag_panel_frame, bg=COLOR_CARD)
        list_container.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 8))
        canvas = tk.Canvas(list_container, bg=COLOR_CARD, highlightthickness=0, height=200)
        scrollbar = ttk.Scrollbar(list_container, orient="vertical", command=canvas.yview)
        self.tag_list_frame = tk.Frame(canvas, bg=COLOR_CARD)
        self.tag_list_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.tag_list_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self._populate_tag_list()

        btn_frame = tk.Frame(self.tag_panel_frame, bg=COLOR_CARD)
        btn_frame.pack(fill=tk.X, padx=8, pady=(0, 10))
        tk.Button(btn_frame, text="➕ 添加标签", command=self._add_tag_dialog,
                  font=ui_font(FS_SMALL, "bold"), bg=COLOR_OK, fg="white",
                  relief=tk.FLAT, borderwidth=0, padx=10, pady=5, cursor="hand2"
                  ).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=2)
        tk.Button(btn_frame, text="🗑️ 删除", command=self._delete_tag,
                  font=ui_font(FS_SMALL, "bold"), bg=COLOR_DANGER, fg="white",
                  relief=tk.FLAT, borderwidth=0, padx=10, pady=5, cursor="hand2"
                  ).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=2)

    def _populate_tag_list(self):
        for w in self.tag_list_frame.winfo_children():
            w.destroy()
        self.tag_check_vars = {}
        for tag, color in self.tag_colors.items():
            frame = tk.Frame(self.tag_list_frame, bg=COLOR_CARD)
            frame.pack(fill=tk.X, pady=2)
            var = tk.BooleanVar(value=False)
            cb = tk.Checkbutton(frame, variable=var, bg=COLOR_CARD, activebackground=COLOR_CARD, selectcolor=COLOR_CARD)
            cb.pack(side=tk.LEFT)
            color_box = tk.Label(frame, text="   ", bg=color, width=4, relief=tk.RIDGE, cursor="hand2")
            color_box.pack(side=tk.LEFT, padx=4)
            color_box.bind("<Button-1>", lambda e, t=tag, c=color: self._on_tag_click(t, c))
            name_label = tk.Label(frame, text=tag, bg=COLOR_CARD, font=ui_font(FS_SMALL), cursor="hand2")
            name_label.pack(side=tk.LEFT, padx=4)
            name_label.bind("<Button-1>", lambda e, t=tag, c=color: self._on_tag_click(t, c))
            self.tag_check_vars[tag] = var

    def _on_tag_click(self, tag_name, tag_color):
        def on_save_and_apply(old_name, new_name, new_color, status_dict):
            # 1. 更新标签颜色与名称（若名称变化则更新所有相关数据）
            if old_name != new_name:
                # 重命名猫咪标签映射中的键
                for cid in list(self.cat_tags_map.keys()):
                    if old_name in self.cat_tags_map[cid]:
                        self.cat_tags_map[cid] = [new_name if t == old_name else t for t in self.cat_tags_map[cid]]
                # 更新标签颜色字典的键
                self.tag_colors.pop(old_name, None)
                self.tag_colors[new_name] = new_color
            else:
                self.tag_colors[old_name] = new_color

            # 保存标签颜色到文件
            save_tag_colors(TAG_COLOR_PATH, self.tag_colors)

            # 2. 根据新的状态字典更新猫咪标签分配（使用新名称）
            for cat_id, has_tag in status_dict.items():
                if cat_id not in self.cat_tags_map:
                    self.cat_tags_map[cat_id] = []
                if has_tag and new_name not in self.cat_tags_map[cat_id]:
                    self.cat_tags_map[cat_id].append(new_name)
                elif not has_tag and new_name in self.cat_tags_map[cat_id]:
                    self.cat_tags_map[cat_id].remove(new_name)

            # 3. 刷新界面与同步 JSON
            self._sync_tags_to_json()
            self._render_tag_panel()

        TagManageWindow(self, tag_name, tag_color, self.rows, self.cat_tags_map,
                        list(self.tag_colors.keys()), on_save_and_apply)

    def _sync_tags_to_json(self):
        regenerate_cats_json(self.rows, CLASSIFIED_DIR, JSON_PATH, self.cat_tags_map)

    def _add_tag_dialog(self):
        dialog = tk.Toplevel(self)
        dialog.title("添加标签")
        sw = self.winfo_screenwidth(); sh = self.winfo_screenheight()
        w, h = int(sw * 0.4), int(sh * 0.5)
        x, y = (sw - w) // 2, (sh - h) // 3
        dialog.geometry(f"{w}x{h}+{x}+{y}")
        dialog.minsize(500, 350)
        dialog.configure(bg=COLOR_BG)
        dialog.transient(self)
        dialog.grab_set()

        tk.Label(dialog, text="标签名称：", bg=COLOR_BG, font=ui_font(FS_BODY)).pack(pady=(20, 4))
        name_var = tk.StringVar()
        tk.Entry(dialog, textvariable=name_var, font=ui_font(FS_BODY), width=20).pack(pady=4)

        tk.Label(dialog, text="选择颜色：", bg=COLOR_BG, font=ui_font(FS_BODY)).pack(pady=(10, 4))
        color_var = tk.StringVar(value="#F9A8D4")
        color_frame = tk.Frame(dialog, bg=COLOR_BG)
        color_frame.pack(pady=6)
        for i, c in enumerate(PRESET_COLORS):
            btn = tk.Button(color_frame, bg=c, width=3, height=1,
                            relief=tk.FLAT, borderwidth=1, cursor="hand2",
                            command=lambda col=c: color_var.set(col))
            btn.grid(row=i // 8, column=i % 8, padx=2, pady=2)
        tk.Button(color_frame, text="🎨", bg="#FFFFFF", font=("", 10), width=3, height=1,
                  relief=tk.FLAT, borderwidth=1, cursor="hand2",
                  command=lambda: self._pick_color(color_var, dialog)).grid(row=len(PRESET_COLORS) // 8, column=len(PRESET_COLORS) % 8, padx=2, pady=2)

        def save():
            name = name_var.get().strip()
            if not name: messagebox.showwarning("提示", "名称不能为空", parent=dialog); return
            if name in self.tag_colors: messagebox.showwarning("提示", "该标签已存在", parent=dialog); return
            self.tag_colors[name] = color_var.get()
            save_tag_colors(TAG_COLOR_PATH, self.tag_colors)
            self._render_tag_panel()
            dialog.destroy()

        tk.Button(dialog, text="保存", command=save, font=ui_font(FS_BODY, "bold"),
                  bg=COLOR_ACCENT, fg="white", padx=20, pady=6, relief=tk.FLAT, cursor="hand2").pack(pady=20)

    def _pick_color(self, color_var, dialog):
        chosen = colorchooser.askcolor(color=color_var.get(), parent=dialog)
        if chosen[1]:
            color_var.set(chosen[1])

    def _delete_tag(self):
        to_delete = [tag for tag, var in self.tag_check_vars.items() if var.get()]
        if not to_delete:
            messagebox.showwarning("提示", "请先勾选要删除的标签"); return
        if not messagebox.askyesno("确认", f"确定删除以下标签吗？\n{', '.join(to_delete)}"): return
        for tag in to_delete:
            self.tag_colors.pop(tag, None)
            for cid in self.cat_tags_map:
                if tag in self.cat_tags_map[cid]:
                    self.cat_tags_map[cid].remove(tag)
        save_tag_colors(TAG_COLOR_PATH, self.tag_colors)
        self._render_tag_panel()
        self._sync_tags_to_json()
        self.status_var.set(f"已删除标签: {', '.join(to_delete)}")

    # ---------- 敏感词管理 ----------
    def _open_ban_words(self):
        BanWordsEditor(self, BAN_WORDS_PATH)

    # ---------- UI 构建 ----------
    def _configure_ttk_styles(self):
        style = ttk.Style(self)
        try: style.theme_use("clam")
        except: pass
        style.configure("TCombobox", font=ui_font(FS_BODY), padding=4)

    def _build_menu(self):
        bar = tk.Menu(self)
        m_file = tk.Menu(bar, tearoff=False)
        m_file.add_command(label="刷新数据", command=self._load_all_data)
        m_file.add_command(label="手动同步前端 (cats.json)", command=self._apply_changes)
        m_file.add_separator()
        m_file.add_command(label="打开 Excel", command=lambda: open_in_explorer(EXCEL_PATH))
        m_file.add_command(label="打开 cats.json", command=lambda: open_in_explorer(JSON_PATH))
        m_file.add_command(label="打开 classified 文件夹", command=lambda: open_in_explorer(CLASSIFIED_DIR))
        m_file.add_separator()
        m_file.add_command(label="退出", command=self._on_close_app)
        bar.add_cascade(label="文件", menu=m_file)
        self.config(menu=bar)

    def _build(self):
        head = tk.Frame(self, bg=COLOR_BG)
        head.pack(fill=tk.X, pady=(20, 4))
        tk.Label(head, text="🐾 SUAT 🐱", font=ui_font(FS_TITLE, "bold"), bg=COLOR_BG, fg=COLOR_ACCENT).pack()
        tk.Label(head, text="记录每一个相遇与告别", font=ui_font(FS_SUBTITLE), bg=COLOR_BG, fg=COLOR_SUBTEXT).pack(pady=(2, 10))

        bar = tk.Frame(self, bg=COLOR_BG)
        bar.pack(fill=tk.X, padx=24, pady=(0, 6))
        def primary(parent, text, color, cmd):
            return tk.Button(parent, text=text, command=cmd, font=ui_font(FS_BODY, "bold"), bg=color, fg="white",
                             relief=tk.FLAT, borderwidth=0, padx=14, pady=7, cursor="hand2")
        primary(bar, "➕ 新增猫咪", COLOR_ACCENT, self._add_cat).pack(side=tk.LEFT, padx=(0, 6))
        # ===敏感词管理按钮===
        primary(bar, "🛡️ 敏感词", "#9B59B6", self._open_ban_words).pack(side=tk.LEFT, padx=4)
        primary(bar, "🔄 刷新", COLOR_INFO, self._load_all_data).pack(side=tk.LEFT, padx=4)
        primary(bar, "💾 应用变更", COLOR_WARN, self._apply_changes).pack(side=tk.LEFT, padx=4)
        self.preview_btn = tk.Button(bar, text="🖥️ 预览前端", font=ui_font(FS_BODY, "bold"), bg=COLOR_OK, fg="white",
                                     relief=tk.FLAT, borderwidth=0, padx=14, pady=7, cursor="hand2", command=self._toggle_server)
        self.preview_btn.pack(side=tk.LEFT, padx=4)

        right_search = tk.Frame(bar, bg=COLOR_BG)
        right_search.pack(side=tk.RIGHT)
        self.var_search = tk.StringVar()
        e = tk.Entry(right_search, textvariable=self.var_search, font=ui_font(FS_BODY), width=22, relief=tk.FLAT,
                     bg="white", highlightthickness=1, highlightbackground=COLOR_DIVIDER, highlightcolor=COLOR_INFO)
        e.pack(side=tk.RIGHT, ipady=5, padx=(4, 0))
        e.bind("<KeyRelease>", lambda _e: self._on_search())
        tk.Label(right_search, text="🔍", bg=COLOR_BG, fg=COLOR_SUBTEXT, font=ui_font(FS_BODY)).pack(side=tk.RIGHT)

        filter_bar = tk.Frame(self, bg=COLOR_BG)
        filter_bar.pack(fill=tk.X, padx=24, pady=(2, 6))
        self.filter_buttons = {}
        for label, key in [("全部猫咪", "all"), ("♂ 男孩", "male"), ("♀ 女孩", "female"),
                            ("❓ 未知", "unknown"), ("⭐ 喵星", "star"), ("🔍 失踪", "lost"), ("🏡 被领养", "adopted")]:
            b = tk.Button(filter_bar, text=label, font=ui_font(FS_SMALL, "bold"), bg="white", fg=COLOR_SUBTEXT,
                          relief=tk.FLAT, borderwidth=0, padx=14, pady=6, cursor="hand2",
                          command=lambda k=key: self._set_filter(k))
            b.pack(side=tk.LEFT, padx=3)
            self.filter_buttons[key] = b
        self._update_filter_buttons()

        main_area = tk.Frame(self, bg=COLOR_BG)
        main_area.pack(fill=tk.BOTH, expand=True, padx=20, pady=(8, 14))

        list_area = tk.Frame(main_area, bg=COLOR_BG)
        list_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.canvas = tk.Canvas(list_area, bg=COLOR_BG, highlightthickness=0)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sb = ttk.Scrollbar(list_area, orient="vertical", command=self.canvas.yview)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.canvas.configure(yscrollcommand=sb.set)
        self.list_frame = tk.Frame(self.canvas, bg=COLOR_BG)
        self._list_window_id = self.canvas.create_window((0, 0), window=self.list_frame, anchor="nw")
        self.list_frame.bind("<Configure>", lambda _e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfig(self._list_window_id, width=e.width))

        right_panel = tk.Frame(main_area, bg=COLOR_CARD, width=300, highlightthickness=1, highlightbackground=COLOR_CARD_BORDER)
        right_panel.pack(side=tk.RIGHT, fill=tk.Y, padx=(12, 0))
        right_panel.pack_propagate(False)
        self.tag_panel_frame = tk.Frame(right_panel, bg=COLOR_CARD)
        self.tag_panel_frame.pack(fill=tk.BOTH, expand=True)

        self.status_var = tk.StringVar(value="就绪")
        status = tk.Label(self, textvariable=self.status_var, bg=COLOR_DIVIDER, fg=COLOR_ACCENT_LIGHT,
                          font=ui_font(FS_SMALL), anchor="w", padx=18, pady=4)
        status.pack(fill=tk.X, side=tk.BOTTOM)

        self._render_tag_panel()

    def _set_filter(self, key):
        self.current_filter = key
        self._update_filter_buttons()
        self._render()

    def _update_filter_buttons(self):
        for k, b in self.filter_buttons.items():
            b.config(bg=COLOR_ACCENT if k == self.current_filter else "white",
                     fg="white" if k == self.current_filter else COLOR_SUBTEXT)

    def _on_search(self):
        self.search_term = self.var_search.get().strip().lower()
        self._render()

    def _add_cat(self):
        editor = CatEditor(self, self.store, cat=None,
                           on_saved=self._on_editor_saved,
                           on_tags_applied=None,
                           tag_colors=list(self.tag_colors.keys()),
                           cat_tags=[])
        self.wait_window(editor)

    def _edit_cat(self, cat):
        cat_tags = self.cat_tags_map.get(cat["id"], [])
        editor = CatEditor(self, self.store, cat=cat,
                           on_saved=self._on_editor_saved,
                           on_tags_applied=self._on_tags_applied_in_editor,
                           tag_colors=list(self.tag_colors.keys()),
                           cat_tags=cat_tags)
        self.wait_window(editor)

    def _on_editor_saved(self, actions, cat_id, tags):
        self.store.save()
        self.cat_tags_map[cat_id] = tags
        self._reload_from_excel()
        count, warns = regenerate_cats_json(self.rows, CLASSIFIED_DIR, JSON_PATH, self.cat_tags_map)
        actions.append(f"重生成 cats.json，共 {count} 只猫")
        msg = "\n".join("• " + a for a in actions)
        if warns: msg += "\n\n⚠️ 警告：\n" + "\n".join(warns[:10])
        messagebox.showinfo("应用成功", msg)
        self.status_var.set("已同步 Excel 与 cats.json")
        self._order_dirty = False

    def _on_tags_applied_in_editor(self, cat_id, tags):
        self.cat_tags_map[cat_id] = tags
        self._sync_tags_to_json()
        self.status_var.set(f"标签已更新：{cat_id}")

    def _reload_from_excel(self):
        try:
            self.store.load()
            self.rows = self.store.all_rows()
            self.thumb_cache.clear()
            self._render()
            self.status_var.set(f"已加载 {len(self.rows)} 只猫咪")
            self._order_dirty = False
        except Exception as e:
            messagebox.showerror("加载失败", str(e))

    def _filtered(self):
        rs = list(self.rows)
        f = self.current_filter
        if f in GENDER_OPTIONS: rs = [r for r in rs if r["gender"] == f]
        elif f in ("star", "lost", "adopted"): rs = [r for r in rs if r["status"] == f]
        if self.search_term:
            t = self.search_term
            rs = [r for r in rs if t in r["name"].lower() or t in r["pic_name"].lower()
                  or t in r["desc"].lower() or t in r["id"]]
        return rs

    def _load_small_thumb(self, cat):
        path = CLASSIFIED_DIR / f"{cat['id']} {cat['name']}" / f"{cat['pic_name']}_01_thumb.jpg"
        if path in self.thumb_cache: return self.thumb_cache[path]
        try:
            if path.is_file(): img = Image.open(path).convert("RGB")
            else: img = Image.new("RGB", (ROW_THUMB_SIZE, ROW_THUMB_SIZE), "#dddddd")
            img.thumbnail((ROW_THUMB_SIZE, ROW_THUMB_SIZE), Image.LANCZOS)
            ph = ImageTk.PhotoImage(img)
        except:
            img = Image.new("RGB", (ROW_THUMB_SIZE, ROW_THUMB_SIZE), "#dddddd")
            ph = ImageTk.PhotoImage(img)
        self.thumb_cache[path] = ph
        return ph

    def _render(self):
        for w in self.list_frame.winfo_children(): w.destroy()
        items = self._filtered()
        if not items:
            tk.Label(self.list_frame, text="没有符合条件的猫咪", bg=COLOR_BG, fg=COLOR_SUBTEXT,
                     font=ui_font(FS_BODY+1), pady=60).pack()
            return
        for i, cat in enumerate(items):
            row_bg = COLOR_CARD if i % 2 == 0 else "#fcfcf9"
            row_frame = tk.Frame(self.list_frame, bg=row_bg, height=64)
            row_frame.pack(fill=tk.X, ipady=4, pady=1)
            thumb = self._load_small_thumb(cat)
            avatar = tk.Label(row_frame, image=thumb, bg=row_bg, cursor="hand2")
            avatar.image = thumb
            avatar.pack(side=tk.LEFT, padx=(8, 12))
            avatar.bind("<Double-Button-1>", lambda _e, c=cat: self._edit_cat(c))

            info_frame = tk.Frame(row_frame, bg=row_bg)
            info_frame.pack(side=tk.LEFT, fill=tk.X, expand=True)
            line1 = tk.Frame(info_frame, bg=row_bg)
            line1.pack(fill=tk.X, anchor="w")
            tk.Label(line1, text=f"{cat['id']} {cat['name']}", bg=row_bg, fg=COLOR_TEXT,
                     font=ui_font(FS_BODY, "bold")).pack(side=tk.LEFT)
            gender_char = {"male": "♂", "female": "♀", "unknown": "?"}[cat["gender"]]
            tk.Label(line1, text=" " + gender_char, bg=row_bg, fg=GENDER_COLOR[cat["gender"]],
                     font=ui_font(FS_BODY, "bold")).pack(side=tk.LEFT)
            status_label = STATUS_LABEL.get(cat["status"], "")
            if status_label:
                tk.Label(line1, text=" " + status_label, bg=row_bg, fg=COLOR_SUBTEXT,
                         font=ui_font(FS_SMALL)).pack(side=tk.LEFT, padx=4)
            line2 = tk.Frame(info_frame, bg=row_bg)
            line2.pack(fill=tk.X, anchor="w")
            n = cat["affection"]
            paws = "🐾" * n + "·" * (5 - n)
            tk.Label(line2, text=paws, bg=row_bg, fg=COLOR_PAW, font=ui_font(FS_SMALL)).pack(side=tk.LEFT)
            desc = (cat["desc"] or "···")[:30]
            tk.Label(line2, text=" " + desc, bg=row_bg, fg=COLOR_SUBTEXT, font=ui_font(FS_SMALL)).pack(side=tk.LEFT)

            btn_frame = tk.Frame(row_frame, bg=row_bg)
            btn_frame.pack(side=tk.RIGHT, padx=4)
            move_var = tk.StringVar(value="1")
            tk.Entry(btn_frame, textvariable=move_var, width=4, font=ui_font(FS_SMALL), relief=tk.FLAT,
                     bg="#f0f0f0", justify="center").pack(side=tk.LEFT, padx=(2, 0))
            for text, cmd in [("▲", "up"), ("▼", "down")]:
                tk.Button(btn_frame, text=text, font=ui_font(FS_SMALL, "bold"), bg=COLOR_HOVER, fg=COLOR_ACCENT,
                          relief=tk.FLAT, padx=6, cursor="hand2",
                          command=lambda c=cat, v=move_var, d=cmd, idx=i: self._move_row(idx, v.get(), d)
                          ).pack(side=tk.LEFT, padx=1)
            tk.Button(btn_frame, text="✏️ 编辑", font=ui_font(FS_SMALL, "bold"), bg=COLOR_HOVER, fg=COLOR_ACCENT,
                      relief=tk.FLAT, padx=8, cursor="hand2", command=lambda c=cat: self._edit_cat(c)
                      ).pack(side=tk.LEFT, padx=(4, 0))
            tk.Button(btn_frame, text="🗑️ 删除", font=ui_font(FS_SMALL, "bold"), bg=COLOR_DANGER, fg="white",
                      relief=tk.FLAT, padx=8, cursor="hand2", command=lambda c=cat: self._delete_cat(c)
                      ).pack(side=tk.LEFT, padx=(4, 0))
            tk.Frame(self.list_frame, bg=COLOR_DIVIDER, height=1).pack(fill=tk.X)

    def _move_row(self, index, step_str, direction):
        try: step = int(step_str)
        except: messagebox.showwarning("输入错误", "请输入整数"); return
        if step < 0: step = -step; direction = "down" if direction == "up" else "up"
        filtered = self._filtered()
        cat = filtered[index]
        real_index = self.rows.index(cat)
        new_index = max(0, real_index - step) if direction == "up" else min(len(self.rows)-1, real_index + step)
        if new_index == real_index: return
        self.rows.pop(real_index)
        self.rows.insert(new_index, cat)
        self.status_var.set(f"已将 {cat['id']} {cat['name']} 移到第 {new_index+1} 位，点击「应用变更」保存")
        self._order_dirty = True
        self._render()

    def _delete_cat(self, cat):
        if not messagebox.askyesno("⚠️ 永久删除", f"确定要删除 {cat['id']} {cat['name']} 吗？"): return
        folder = CLASSIFIED_DIR / f"{cat['id']} {cat['name']}"
        if folder.exists(): shutil.rmtree(folder)
        row_idx = cat.get("_row")
        if row_idx: self.store.delete_row(row_idx); self.store.save()
        self.cat_tags_map.pop(cat["id"], None)
        self._reload_from_excel()
        regenerate_cats_json(self.rows, CLASSIFIED_DIR, JSON_PATH, self.cat_tags_map)
        self.status_var.set(f"已删除 {cat['id']} {cat['name']}")

    def _toggle_server(self):
        if self.http_server: self._stop_server(); return
        try:
            port = 8800
            import socket
            while True:
                with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
                    if s.connect_ex(('localhost', port)) != 0: break
                    port += 1
            class QuietHandler(SimpleHTTPRequestHandler):
                def __init__(self, *args, **kwargs): super().__init__(*args, directory=str(BASE_DIR), **kwargs)
                def log_message(self, format, *args): pass
            self.http_server = HTTPServer(('localhost', port), QuietHandler)
            self.http_server_thread = threading.Thread(target=self.http_server.serve_forever, daemon=True)
            self.http_server_thread.start()
            webbrowser.open(f"http://localhost:{port}/index.html")
            self.preview_btn.config(text="🛑 停止服务器", bg=COLOR_DANGER)
            self.status_var.set(f"预览服务器运行中：http://localhost:{port}/index.html")
        except Exception as e:
            messagebox.showerror("启动失败", str(e))

    def _stop_server(self):
        if self.http_server:
            self.http_server.shutdown()
            self.http_server = None
            self.preview_btn.config(text="🖥️ 预览前端", bg=COLOR_OK)
            self.status_var.set("预览服务器已停止")


# ============================================================
# 入口
# ============================================================
def main():
    if not EXCEL_PATH.exists():
        print(f"找不到 {EXCEL_PATH}")
        sys.exit(1)
    if not CLASSIFIED_DIR.exists():
        CLASSIFIED_DIR.mkdir(parents=True, exist_ok=True)
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()